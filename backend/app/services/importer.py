"""后台数据批量导入解析与落库。"""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import AdmissionStat, Institution, InstitutionMajor, Major


def parse_csv(content: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(content))
    return [dict(row) for row in reader]


def parse_excel(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel 导入需要安装 openpyxl") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [str(cell) if cell is not None else "" for cell in next(rows_iter)]
    result: list[dict[str, Any]] = []
    for row in rows_iter:
        result.append(
            {
                headers[index]: value
                for index, value in enumerate(row)
                if index < len(headers)
            }
        )
    workbook.close()
    return result


def _bool(value: Any, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "是", "yes"}


def _int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _json(value: Any, default: Any) -> Any:
    if value is None or value == "":
        return default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return default


_EMPTY_MARKERS = {"", "暂无", "无", "不公布", "未知", "n/a", "na", "none"}


def _optional_int(value: Any) -> int | None:
    """将空值/“暂无”解析为 None；其余必须是非负整数，否则报错。"""
    if value is None:
        return None
    text = str(value).strip().lower()
    if text in _EMPTY_MARKERS:
        return None
    try:
        number = float(text)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"数值字段无法解析为整数：{value!r}") from exc
    if not number.is_integer() or number < 0:
        raise ValueError(f"数值字段必须为非负整数：{value!r}")
    return int(number)


def _optional_float(value: Any) -> float | None:
    """将空值/“暂无”解析为 None；其余必须是可解析数值。"""
    if value is None:
        return None
    text = str(value).strip().lower()
    if text in _EMPTY_MARKERS:
        return None
    try:
        return float(text)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"数值字段无法解析：{value!r}") from exc


def _resolve_institution(db: Session, row: dict[str, Any]) -> Institution | None:
    if row.get("institution_id"):
        return db.get(Institution, row["institution_id"])
    code = row.get("institution_code")
    if code:
        return db.scalar(select(Institution).where(Institution.code == str(code)))
    return None


def _resolve_major(db: Session, row: dict[str, Any]) -> Major | None:
    if row.get("major_id"):
        return db.get(Major, row["major_id"])
    code = row.get("major_code")
    if code:
        return db.scalar(select(Major).where(Major.code == str(code)))
    return None


def _resolve_institution_major(
    db: Session, row: dict[str, Any]
) -> InstitutionMajor | None:
    if row.get("institution_major_id"):
        return db.get(InstitutionMajor, row["institution_major_id"])
    institution = _resolve_institution(db, row)
    major = _resolve_major(db, row)
    if institution is None or major is None:
        return None
    return db.scalar(
        select(InstitutionMajor).where(
            InstitutionMajor.institution_id == institution.id,
            InstitutionMajor.major_id == major.id,
        )
    )


def import_rows(db: Session, entity_type: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    if entity_type == "institution":
        return _import_institutions(db, rows)
    if entity_type == "major":
        return _import_majors(db, rows)
    if entity_type == "institution_major":
        return _import_institution_majors(db, rows)
    if entity_type == "admission_stat":
        return _import_admission_stats(db, rows)
    raise ValueError(f"不支持的实体类型：{entity_type}")


def _import_institutions(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    created = 0
    updated = 0
    for row in rows:
        code = str(row.get("code", "")).strip()
        if not code or not row.get("name"):
            continue
        item = db.scalar(select(Institution).where(Institution.code == code))
        data = {
            "name": str(row["name"]),
            "short_name": str(row.get("short_name", "")),
            "province": str(row.get("province", "")),
            "city": str(row.get("city", "")),
            "level": str(row.get("level", "普通")),
            "category": str(row.get("category", "综合")),
            "is_self_draw": _bool(row.get("is_self_draw")),
            "official_url": row.get("official_url") or None,
            "logo_url": row.get("logo_url") or None,
            "tags": _json(row.get("tags"), []),
            "status": "published",
        }
        if item is None:
            db.add(Institution(code=code, **data))
            created += 1
        else:
            for key, value in data.items():
                setattr(item, key, value)
            updated += 1
    db.commit()
    return {"entity_type": "institution", "created": created, "updated": updated}


def _import_majors(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    created = 0
    updated = 0
    for row in rows:
        code = str(row.get("code", "")).strip()
        if not code or not row.get("name"):
            continue
        item = db.scalar(select(Major).where(Major.code == code))
        data = {
            "name": str(row["name"]),
            "discipline_gate": str(row.get("discipline_gate", "")),
            "discipline_level1": str(row.get("discipline_level1", "")),
            "degree_type": str(row.get("degree_type", "学硕")),
            "is_cross_allowed": _bool(row.get("is_cross_allowed"), True),
            "cross_condition": row.get("cross_condition") or None,
            "apply_condition": row.get("apply_condition") or None,
            "research_directions": _json(row.get("research_directions"), []),
            "status": "published",
        }
        if item is None:
            db.add(Major(code=code, **data))
            created += 1
        else:
            for key, value in data.items():
                setattr(item, key, value)
            updated += 1
    db.commit()
    return {"entity_type": "major", "created": created, "updated": updated}


def _import_institution_majors(
    db: Session, rows: list[dict[str, Any]]
) -> dict[str, Any]:
    created = 0
    updated = 0
    for row in rows:
        institution = _resolve_institution(db, row)
        major = _resolve_major(db, row)
        if institution is None or major is None:
            continue
        item = db.scalar(
            select(InstitutionMajor).where(
                InstitutionMajor.institution_id == institution.id,
                InstitutionMajor.major_id == major.id,
            )
        )
        data = {
            "faculty_name": str(row.get("faculty_name", "")),
            "study_mode": str(row.get("study_mode", "全日制")),
            "length": _float(row.get("length")),
            "tuition": row.get("tuition") or None,
            "scholarship": row.get("scholarship") or None,
            "exam_basic": _json(row.get("exam_basic"), {}),
            "exam_major": _json(row.get("exam_major"), {}),
            "reexam_form": _json(row.get("reexam_form"), {}),
            "additional_exam": row.get("additional_exam") or None,
            "status": "published",
        }
        if item is None:
            db.add(
                InstitutionMajor(
                    institution_id=institution.id,
                    major_id=major.id,
                    **data,
                )
            )
            created += 1
        else:
            for key, value in data.items():
                setattr(item, key, value)
            updated += 1
    db.commit()
    return {"entity_type": "institution_major", "created": created, "updated": updated}


def _validate_stat_row(
    db: Session,
    row: dict[str, Any],
    index: int,
) -> tuple[dict[str, Any] | None, list[str]]:
    """解析并校验一条招录数据；返回 (可写入的数据, 错误列表)。"""
    errors: list[str] = []
    prefix = f"第 {index} 行"

    institution_major = _resolve_institution_major(db, row)
    if institution_major is None:
        return None, [f"{prefix}: 未找到对应的院校-专业组合"]

    year = _optional_int(row.get("year"))
    if year is None or not 2000 <= year <= 2100:
        errors.append(f"{prefix}: year 缺失或超出 2000-2100")

    data_quality = str(row.get("data_quality") or "official").strip()
    allowed_quality = {"official", "mock", "none", "third_party", "user_submitted"}
    if data_quality not in allowed_quality:
        errors.append(f"{prefix}: data_quality 非法值 {data_quality!r}")

    source_url = (row.get("source_url") or "").strip() or None
    source_name = str(row.get("source_name") or "").strip()
    source_year = _optional_int(row.get("source_year")) or year
    if data_quality == "official":
        if not source_url or not source_url.lower().startswith(("http://", "https://")):
            errors.append(f"{prefix}: official 数据必须提供有效 source_url")
        if not source_name:
            errors.append(f"{prefix}: official 数据必须提供 source_name")
        if source_year != year:
            errors.append(f"{prefix}: official 数据 source_year 必须与 year 一致")

    def get_int(field: str) -> int | None:
        try:
            return _optional_int(row.get(field))
        except ValueError as exc:
            errors.append(f"{prefix}: {field} {exc}")
            return None

    def get_float(field: str) -> float | None:
        try:
            return _optional_float(row.get(field))
        except ValueError as exc:
            errors.append(f"{prefix}: {field} {exc}")
            return None

    plan_total = get_int("plan_total")
    plan_unified = get_int("plan_unified")
    recommended_count = get_int("recommended_count")
    recommended_ratio = get_float("recommended_ratio")
    applicant_count = get_int("applicant_count")
    admitted_count = get_int("admitted_count")
    report_rate = get_float("report_rate")
    reexam_count = get_int("reexam_count")
    reexam_admit_rate = get_float("reexam_admit_rate")
    max_score = get_float("max_score")
    min_score = get_float("min_score")
    avg_score = get_float("avg_score")
    national_line = get_float("national_line")
    self_line = get_float("self_line")
    college_line = get_float("college_line")
    transfer_quota = get_int("transfer_quota")

    if plan_unified is not None and plan_total is not None and plan_unified > plan_total:
        errors.append(f"{prefix}: plan_unified 不能大于 plan_total")
    if recommended_count is not None and plan_total is not None and recommended_count > plan_total:
        errors.append(f"{prefix}: recommended_count 不能大于 plan_total")
    if (
        recommended_ratio is not None
        and not 0 <= recommended_ratio <= 1
    ):
        errors.append(f"{prefix}: recommended_ratio 应为 0-1 的小数比例")
    for field, value in (
        ("report_rate", report_rate),
        ("reexam_admit_rate", reexam_admit_rate),
    ):
        if value is not None and value < 0:
            errors.append(f"{prefix}: {field} 不能为负数")

    scores = {
        "max_score": max_score,
        "min_score": min_score,
        "avg_score": avg_score,
        "national_line": national_line,
        "self_line": self_line,
        "college_line": college_line,
    }
    for field, value in scores.items():
        if value is not None and not 0 <= value <= 500:
            errors.append(f"{prefix}: {field} 超出合理范围 0-500")
    if None not in (min_score, avg_score, max_score):
        if not (min_score <= avg_score <= max_score):
            errors.append(f"{prefix}: 分数逻辑不成立 min_score <= avg_score <= max_score")
    if (
        institution_major.institution.is_self_draw
        and self_line is not None
        and college_line is not None
        and college_line < self_line
    ):
        errors.append(f"{prefix}: 自划线院校 college_line 不应低于 self_line")

    if errors or year is None:
        return None, errors

    data = {
        "plan_total": plan_total,
        "plan_unified": plan_unified,
        "recommended_count": recommended_count,
        "recommended_ratio": recommended_ratio,
        "applicant_count": applicant_count,
        "admitted_count": admitted_count,
        "report_rate": report_rate,
        "reexam_count": reexam_count,
        "reexam_admit_rate": reexam_admit_rate,
        "max_score": max_score,
        "min_score": min_score,
        "avg_score": avg_score,
        "national_line": national_line,
        "self_line": self_line,
        "college_line": college_line,
        "transfer_quota": transfer_quota,
        "metrics": _json(row.get("metrics"), {}),
        "source_url": source_url,
        "source_name": source_name or "人工导入",
        "source_year": source_year or year,
        "data_quality": data_quality,
        "review_status": str(row.get("review_status") or "pending"),
    }
    return data, errors


def _import_admission_stats(
    db: Session, rows: list[dict[str, Any]]
) -> dict[str, Any]:
    created = 0
    updated = 0
    errors: list[str] = []
    for index, row in enumerate(rows, start=1):
        data, row_errors = _validate_stat_row(db, row, index)
        if row_errors:
            errors.extend(row_errors)
            continue
        assert data is not None
        institution_major = _resolve_institution_major(db, row)
        assert institution_major is not None
        year = _optional_int(row.get("year"))
        assert year is not None
        item = db.scalar(
            select(AdmissionStat).where(
                AdmissionStat.institution_major_id == institution_major.id,
                AdmissionStat.year == year,
            )
        )
        if item is None:
            db.add(
                AdmissionStat(
                    institution_major_id=institution_major.id,
                    year=year,
                    **data,
                )
            )
            created += 1
        else:
            for key, value in data.items():
                setattr(item, key, value)
            updated += 1
    db.commit()
    return {
        "entity_type": "admission_stat",
        "created": created,
        "updated": updated,
        "errors": errors,
    }
